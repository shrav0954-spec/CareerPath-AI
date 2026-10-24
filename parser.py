import io
import os
from fastapi import UploadFile, HTTPException
from pypdf import PdfReader
from docx import Document
from app.config import settings

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extract text content from PDF file bytes using pypdf.
    """
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        extracted_text = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                extracted_text.append(page_text)
        return "\n".join(extracted_text)
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Failed to read PDF document: {str(e)}"
        )

def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract text content from Microsoft Word .docx file bytes.
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table text if any
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    paragraphs.append(row_text)
        return "\n".join(paragraphs)
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Failed to read DOCX document: {str(e)}"
        )

def extract_text_from_txt(file_bytes: bytes) -> str:
    """
    Decode plain text files with UTF-8, falling back to latin-1.
    """
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1", errors="ignore")

async def parse_resume_file(file: UploadFile) -> str:
    """
    Validates uploaded resume file extension and size,
    then extracts and sanitizes clean text content.
    """
    filename = file.filename or ""
    _, ext = os.path.splitext(filename.lower())
    
    if ext not in settings.ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type '{ext}'. Please upload a .pdf, .docx, or .txt file."
        )

    # Read file content into memory
    content = await file.read()
    
    # Check file size limit
    size_mb = len(content) / (1024 * 1024)
    if size_mb > settings.MAX_FILE_SIZE_MB:
        raise HTTPException(
            status_code=400,
            detail=f"File exceeds maximum size of {settings.MAX_FILE_SIZE_MB}MB (received {size_mb:.1f}MB)."
        )

    # Parse depending on format
    if ext == ".pdf":
        raw_text = extract_text_from_pdf(content)
    elif ext == ".docx":
        raw_text = extract_text_from_docx(content)
    elif ext == ".txt":
        raw_text = extract_text_from_txt(content)
    else:
        raise HTTPException(status_code=400, detail="Unsupported file format.")

    # Clean whitespace and strip
    cleaned_text = "\n".join([line.strip() for line in raw_text.splitlines() if line.strip()])

    if len(cleaned_text.strip()) < 50:
        raise HTTPException(
            status_code=400,
            detail="The uploaded resume seems to be empty or contains scanned images without selectable text. Please upload a text-based resume."
        )

    # Truncate extremely long documents (over 30,000 chars) to prevent token overflow
    max_chars = 30000
    if len(cleaned_text) > max_chars:
        cleaned_text = cleaned_text[:max_chars]

    return cleaned_text
